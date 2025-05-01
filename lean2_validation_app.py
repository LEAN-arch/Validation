# lean2_validation_app.py

import streamlit as st
import pandas as pd
import io
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows

st.set_page_config(page_title="LEAN 2.0 Validation Toolkit", layout="wide")

# App state
if "metrics" not in st.session_state:
    st.session_state.metrics = []
if "feedback" not in st.session_state:
    st.session_state.feedback = []
if "audit" not in st.session_state:
    st.session_state.audit = {}

# Sidebar: Navigation
st.sidebar.image("assets/FOBO2.png", width=120)
st.sidebar.title("LEAN 2.0 Validation Toolkit")
section = st.sidebar.radio(
    "Navigate",
    ["Pilot Checklist", "Metrics Dashboard", "Feedback Log", "Ethics Audit", "Download Report"],
    format_func=lambda x: f"▶ {x}"
)

# CSS Styling
st.markdown("""
<style>
.big-font { font-size:18px !important; }
.section-title { font-size:26px; font-weight:600; margin-top:20px; }
</style>
""", unsafe_allow_html=True)

# Excel export

def export_to_excel(metric_df=None, feedback_list=None, audit_dict=None):
    wb = Workbook()
    ws_metrics = wb.active
    ws_metrics.title = "Metrics"

    if metric_df is not None:
        for r in dataframe_to_rows(metric_df, index=False, header=True):
            ws_metrics.append(r)

    if feedback_list:
        ws_feedback = wb.create_sheet("Stakeholder Feedback")
        ws_feedback.append(["Role", "Name", "Observations", "Concerns", "Suggestions"])
        for entry in feedback_list:
            ws_feedback.append([
                entry.get("stakeholder"), entry.get("name"), entry.get("observations"),
                entry.get("concerns"), entry.get("suggestions")
            ])

    if audit_dict:
        ws_audit = wb.create_sheet("Ethics Audit")
        ws_audit.append(["Audit Item", "Result"])
        for item, result in audit_dict.items():
            ws_audit.append([item, result])

    excel_io = io.BytesIO()
    wb.save(excel_io)
    excel_io.seek(0)
    return excel_io

# Section: Pilot Checklist
if section == "Pilot Checklist":
    st.markdown('<div class="section-title">Pilot Readiness Checklist</div>', unsafe_allow_html=True)
    st.info("Ensure the pilot is ethically aligned and operationally scoped before launching.")

    with st.form("checklist_form"):
        col1, col2, col3 = st.columns(3)
        with col1:
            st.subheader("Pilot Scope")
            pilot_area = st.checkbox("Area selected with baseline")
            goals_defined = st.checkbox("Goals: Efficiency + Well-being")
            control_established = st.checkbox("Control/Baseline defined")
            team_involved = st.checkbox("Cross-functional team ready")

        with col2:
            st.subheader("Tool Setup")
            tools_selected = st.checkbox("LEAN 2.0 tools identified")
            facilitators_trained = st.checkbox("Ethical facilitators trained")
            tech_integrated = st.checkbox("Measurement tools connected")

        with col3:
            st.subheader("Measurement Plan")
            kpis_selected = st.checkbox("KPI mix: Ops, Safety, Ethics")
            timeline_set = st.checkbox("Pilot timeline finalized")

        if st.form_submit_button("✔ Submit Checklist"):
            st.success("Checklist submitted successfully!")

# Section: Metrics
elif section == "Metrics Dashboard":
    st.markdown('<div class="section-title">Metrics Tracking</div>', unsafe_allow_html=True)
    st.caption("Track operational + ethical impact of LEAN 2.0 pilot over time.")

    with st.form("metrics_form"):
        col1, col2 = st.columns([3, 1])
        with col1:
            name = st.text_input("Metric Name", placeholder="e.g., Worker Satisfaction")
        with col2:
            mtype = st.selectbox("Category", ["Operational", "Well-being", "Safety", "Ethics"])

        pre = st.number_input("Pre-Pilot Value", step=0.1)
        wk2 = st.number_input("Week 2 Value", step=0.1)
        wk4 = st.number_input("Week 4 Value", step=0.1)

        if st.form_submit_button("Add Metric"):
            change = round(((wk4 - pre) / pre) * 100, 2) if pre else 0
            st.session_state.metrics.append({
                "Metric": name, "Type": mtype, "Pre": pre,
                "Week 2": wk2, "Week 4": wk4, "Δ (%)": change
            })
            st.success(f"Added metric '{name}'.")

    if st.session_state.metrics:
        st.dataframe(pd.DataFrame(st.session_state.metrics), use_container_width=True)

# Section: Feedback
elif section == "Feedback Log":
    st.markdown('<div class="section-title">Stakeholder Feedback</div>', unsafe_allow_html=True)
    st.caption("Capture diverse voices: frontline, quality, ops, HR, etc.")

    with st.form("feedback_form"):
        col1, col2 = st.columns(2)
        with col1:
            stakeholder = st.text_input("Stakeholder Role", placeholder="e.g., Operator, Engineer")
        with col2:
            name = st.text_input("Name (optional)")

        observations = st.text_area("Observations (what changed?)")
        concerns = st.text_area("Concerns or risks")
        suggestions = st.text_area("Improvement suggestions")

        if st.form_submit_button("Submit Feedback"):
            st.session_state.feedback.append({
                "stakeholder": stakeholder,
                "name": name,
                "observations": observations,
                "concerns": concerns,
                "suggestions": suggestions
            })
            st.success("Feedback saved.")

    if st.session_state.feedback:
        st.write(pd.DataFrame(st.session_state.feedback))

# Section: Ethics Audit
elif section == "Ethics Audit":
    st.markdown('<div class="section-title">Ethics Alignment Audit</div>', unsafe_allow_html=True)
    st.caption("Review ethical compliance for fairness, autonomy, and safety.")

    questions = {
        "Workload Balance": "Are takt times adjusted for human sustainability?",
        "Worker Autonomy": "Do workers have Kaizen input authority?",
        "Psych Safety": "Are dissenting voices protected?",
        "Overburden Detection": "Is fatigue/stress actively monitored?",
        "Fairness": "Are all groups represented in process design?",
        "Feedback Transparency": "Are decisions based on open feedback?"
    }

    with st.form("audit_form"):
        results = {}
        for key, q in questions.items():
            results[key] = st.radio(q, ["Pass", "Fail"], horizontal=True, key=key)
        if st.form_submit_button("Finalize Audit"):
            st.session_state.audit = results
            st.success("Ethics audit results recorded.")

    if st.session_state.audit:
        st.json(st.session_state.audit)

# Section: Download
elif section == "Download Report":
    st.markdown('<div class="section-title">Download Validation Report</div>', unsafe_allow_html=True)
    metrics_df = pd.DataFrame(st.session_state.metrics) if st.session_state.metrics else None
    feedback_list = st.session_state.feedback
    audit_dict = st.session_state.audit

    if metrics_df is not None or feedback_list or audit_dict:
        excel_file = export_to_excel(metrics_df, feedback_list, audit_dict)
        st.download_button("📅 Download Excel Report",
                           data=excel_file,
                           file_name="LEAN2_Validation_Report.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        st.warning("No validation data entered yet.")
